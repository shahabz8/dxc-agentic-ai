"""Ingestion: load AskIT knowledge-base articles (markdown + front-matter metadata), then chunk them."""
import glob, io, os, re, time
from datetime import date
from dataclasses import dataclass, field


@dataclass
class Section:
    n: str
    heading: str
    text: str


@dataclass
class Document:
    doc_id: str
    meta: dict
    sections: list = field(default_factory=list)

    @property
    def title(self):
        return self.meta.get("title", self.doc_id)


@dataclass
class Chunk:
    chunk_id: str
    text: str
    doc: Document
    section: Section | None = None

    @property
    def citation(self):
        m = self.doc.meta
        base = f"{m['code']} {m['version']}"
        return f"{base} §{self.section.n}" if self.section else f"{base} ({self.chunk_id})"

    def meta_row(self):
        m = self.doc.meta
        return dict(article=m["code"], version=m["version"], audience=m["audience"],
                    effective=m["effective"], status=m["status"])


def parse_markdown(raw, doc_id, default_meta=None):
    meta = dict(default_meta or {})
    raw = raw.replace("\r\n", "\n")
    m = re.match(r"^---\n(.*?)\n---\n(.*)$", raw, re.S)
    body = raw
    if m:
        for line in m.group(1).splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip()
        body = m.group(2)
    doc = Document(doc_id, meta)
    blocks = re.split(r"^#{1,3} ", body, flags=re.M)
    if len(blocks) > 1:
        for i, block in enumerate(blocks[1:], 1):
            head, _, text = block.partition("\n")
            num, dot, h = head.partition(". ")
            if not (dot and num.strip().isdigit()):
                num, h = str(i), head
            text = " ".join(text.split())
            if text:
                doc.sections.append(Section(num.strip(), h.strip(), text))
    else:
        doc.sections = _group_paragraphs(re.split(r"\n\s*\n", body))
    return doc


def load_documents(folder=None):
    folder = folder or os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "kb")
    docs = []
    for path in sorted(glob.glob(os.path.join(folder, "*.md"))):
        raw = open(path, encoding="utf-8").read()
        docs.append(parse_markdown(raw, os.path.splitext(os.path.basename(path))[0]))
    return docs


def _group_paragraphs(paras, max_chars=800, label="Part"):
    """No headings found: group paragraphs into sections of up to ~max_chars."""
    secs, buf = [], ""
    for p in (" ".join(x.split()) for x in paras):
        if not p:
            continue
        if buf and len(buf) + len(p) > max_chars:
            secs.append(buf); buf = ""
        buf = f"{buf} {p}".strip()
    if buf:
        secs.append(buf)
    return [Section(str(i), f"{label} {i}", t) for i, t in enumerate(secs, 1)]


INJECTION_PATTERNS = [r"ignore (all )?(the )?(previous|prior|above) (instructions|rules)", r"ignore all previous",
                      r"note for the (ai|assistant)", r"system (note|prompt|override)", r"do not mention this",
                      r"send (us )?(your|their) (current )?password"]


def scan_injection(doc):
    """Very simple prompt-injection scan: returns the section headings whose text looks like instructions to the AI.
    It only WARNS. Real defences are layered (see the lab README)."""
    hits = []
    for s in doc.sections:
        if any(re.search(p, s.text, re.I) for p in INJECTION_PATTERNS):
            hits.append(s.heading)
    return hits


MAX_UPLOAD_MB = 5
MAX_PDF_PAGES = 60


def parse_upload(name, data, meta):
    """Parse an uploaded .md/.txt/.pdf/.docx into a Document. Returns (doc, info)."""
    t0 = time.perf_counter()
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", os.path.splitext(name)[0])[:40] or "upload"
    ext = name.lower().rsplit(".", 1)[-1]
    base = dict(code=stem, title=os.path.splitext(name)[0], version="v1", audience="All",
                effective=date.today().isoformat(), status="active")
    base.update({k: v for k, v in meta.items() if v})
    info = dict(file=name, type=ext.upper(), size_kb=round(len(data) / 1024, 1), pages=None, warning="")
    if len(data) > MAX_UPLOAD_MB * 1024 * 1024:
        raise ValueError(f"{name} is larger than {MAX_UPLOAD_MB} MB")
    if ext in ("md", "txt"):
        doc = parse_markdown(data.decode("utf-8", errors="ignore"), "up_" + stem, base)
        doc.meta.update({k: v for k, v in meta.items() if v})
    elif ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        info["pages"] = len(reader.pages)
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError(f"{name} has {len(reader.pages)} pages; the demo limit is {MAX_PDF_PAGES}")
        doc = Document("up_" + stem, base)
        head_re = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+([A-Z][^.]{1,70})$")
        n, cur_h, buf = 0, None, []
        def flush(page):
            nonlocal n
            text = " ".join(" ".join(buf).split())
            if text:
                n += 1
                doc.sections.append(Section(str(n), cur_h or f"Page {page}", text))
        for i, page in enumerate(reader.pages, 1):
            for line in (page.extract_text() or "").splitlines():
                m = head_re.match(line.strip())
                if m:  # numbered heading like "2. Eligible programs" -> start a new section
                    flush(i); buf = []; cur_h = m.group(2).strip()
                elif line.strip():
                    buf.append(line.strip())
            if cur_h is None:  # no headings detected: one section per page
                flush(i); buf = []
        flush(len(reader.pages))
        if not doc.sections:
            info["warning"] = "No text layer found. Probably a scanned PDF: needs OCR."
    elif ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        doc, cur_h, buf, n = Document("up_" + stem, base), None, [], 0
        def flush():
            nonlocal n
            text = " ".join(" ".join(buf).split())
            if text:
                n += 1
                doc.sections.append(Section(str(n), cur_h or f"Part {n}", text))
        for p in d.paragraphs:
            if p.style is not None and p.style.name.lower().startswith(("heading", "title")) and p.text.strip():
                flush(); buf = []; cur_h = p.text.strip()
            elif p.text.strip():
                buf.append(p.text.strip())
        flush()
        for t in d.tables:
            rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
            n += 1
            doc.sections.append(Section(str(n), f"Table {n}", " ; ".join(rows)))
        if len(doc.sections) <= 1 and doc.sections:
            doc.sections = _group_paragraphs([p.text for p in d.paragraphs])
    else:
        raise ValueError(f"{name}: unsupported type .{ext}")
    info.update(sections=len(doc.sections), chars=sum(len(s.text) for s in doc.sections),
                parse_ms=int((time.perf_counter() - t0) * 1000))
    return doc, info


def chunk_documents(docs, strategy="section", size=180, overlap=0):
    chunks = []
    if strategy == "section":
        for d in docs:
            for s in d.sections:
                if len(s.text) <= 1200:
                    chunks.append(Chunk(f"{d.doc_id}#s{s.n}", f"{s.heading}. {s.text}", d, s))
                    continue
                # long section (e.g. a PDF page): split at sentence boundaries into ~800-char parts
                parts, buf = [], ""
                for sent in re.split(r"(?<=[.!?])\s+", s.text):
                    if buf and len(buf) + len(sent) > 800:
                        parts.append(buf); buf = ""
                    buf = f"{buf} {sent}".strip()
                if buf:
                    parts.append(buf)
                for k, part in enumerate(parts, 1):
                    chunks.append(Chunk(f"{d.doc_id}#s{s.n}.{k}", f"{s.heading}. {part}", d, s))
        return chunks
    # fixed-size character windows (with optional overlap) on the flattened text
    step = max(20, size - overlap)
    for d in docs:
        full = f"{d.title}. " + " ".join(s.text for s in d.sections)
        i, k = 0, 1
        while i < len(full):
            end = min(len(full), i + size)
            if end < len(full):
                sp = full.rfind(" ", i, end)
                end = sp if sp > i + size // 2 else end
            chunks.append(Chunk(f"{d.doc_id}#c{k}", full[i:end].strip(), d, None))
            if end >= len(full):
                break
            i = max(end - overlap, i + 1)
            k += 1
    return chunks
